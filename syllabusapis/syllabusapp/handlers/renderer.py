import io
import json
import os

from django.conf import settings
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

from handlers.base import TextRenderer, TableRenderer, set_table_borders, SelectionRenderer, ReferenceRenderer
from syllabuses.models import Syllabus


def int_to_roman(num):
    val = [1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1]
    syb = ["M", "CM", "D", "CD", "C", "XC", "L", "XL", "X", "IX", "V", "IV", "I"]
    roman_num = ''
    i = 0
    while num > 0:
        for _ in range(num // val[i]):
            roman_num += syb[i]
            num -= val[i]
        i += 1
    return roman_num

class SyllabusDocxRenderer:
    def __init__(self, syllabus_id):
        self.syllabus = Syllabus.objects.prefetch_related(
            'main_sections__sub_sections'
        ).get(id=syllabus_id)
        template_path = os.path.join(
            settings.BASE_DIR,
            'templates', 'exports', 'styles_blank.docx'
        )
        self.doc = Document(template_path)

        style = self.doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(13)

        self.renderers = {
            'text': TextRenderer(self.doc),
            'table': TableRenderer(self.doc),
            'reference': ReferenceRenderer(self.doc),
            'selection': SelectionRenderer(self.doc)
        }

    def _render_document_header(self):

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        table.autofit = False
        table.columns[0].width = Cm(4.5)
        table.cell(0, 0).width = Cm(4.5)
        set_table_borders(table)

        p_phuluc = table.cell(0, 0).paragraphs[0]
        p_phuluc.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_phuluc = p_phuluc.add_run("Phụ lục 7-Appendix 7")
        run_phuluc.font.name = 'Times New Roman'
        run_phuluc.font.italic = True
        run_phuluc.font.bold = True
        run_phuluc.font.size = Pt(12)

        self.doc.add_paragraph()

        p_school = self.doc.add_paragraph()
        p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_bo = p_school.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nMINISTRY OF EDUCATION AND TRAINING\n")
        run_bo.font.name = 'Times New Roman'
        run_bo.font.size = Pt(13)

        run_truong = p_school.add_run("TRƯỜNG ĐẠI HỌC MỞ THÀNH PHỐ HỒ CHÍ MINH\nHO CHI MINH CITY OPEN UNIVERSITY\n")
        run_truong.font.name = 'Times New Roman'
        run_truong.font.bold = True
        run_truong.font.size = Pt(13)

        run_line = p_school.add_run("______________________________")
        run_line.font.name = 'Times New Roman'
        run_line.font.bold = True

        self.doc.add_paragraph()

        p_title = self.doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_title_vi = p_title.add_run("ĐỀ CƯƠNG MÔN HỌC\n")
        run_title_vi.font.name = 'Times New Roman'
        run_title_vi.font.bold = True
        run_title_vi.font.size = Pt(16)

        run_title_en = p_title.add_run("COURSE SPECIFICATION\n")
        run_title_en.font.name = 'Times New Roman'
        run_title_en.font.bold = True
        run_title_en.font.size = Pt(14)

    def _render_document_footer(self):
        self.doc.add_paragraph()

        p_date = self.doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_date = p_date.add_run("Thành phố Hồ Chí Minh, ngày ... tháng ... năm 202...")
        run_date.font.name = 'Times New Roman'
        run_date.font.italic = True
        run_date.font.size = Pt(13)

        table = self.doc.add_table(rows=1, cols=2)

        cell_left = table.cell(0, 0)
        p_left = cell_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_left_title = p_left.add_run("PHỤ TRÁCH KHOA CNTT\nDEPUTY DEAN\n")
        run_left_title.font.name = 'Times New Roman'
        run_left_title.font.bold = True
        run_left_title.font.size = Pt(13)

        run_left_note = p_left.add_run("(Ký và ghi rõ họ tên-Signed with fullname)\n\n\n\n\n")
        run_left_note.font.name = 'Times New Roman'
        run_left_note.font.italic = True
        run_left_note.font.size = Pt(13)

        run_left_name = p_left.add_run("TS. Trương Hoàng Vinh")
        run_left_name.font.name = 'Times New Roman'
        run_left_name.font.bold = True
        run_left_name.font.size = Pt(13)

        cell_right = table.cell(0, 1)
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run_right_title = p_right.add_run("GIẢNG VIÊN BIÊN SOẠN\nACADEMIC\n")
        run_right_title.font.name = 'Times New Roman'
        run_right_title.font.bold = True
        run_right_title.font.size = Pt(13)

        run_right_note = p_right.add_run("(Ký và ghi rõ họ tên- Signed with fullname)\n\n\n\n\n")
        run_right_note.font.name = 'Times New Roman'
        run_right_note.font.italic = True
        run_right_note.font.size = Pt(13)

        run_right_name = p_right.add_run("ThS. Dương Hữu Thành")
        run_right_name.font.name = 'Times New Roman'
        run_right_name.font.bold = True
        run_right_name.font.size = Pt(13)

    def render(self):
        self._render_document_header()

        for main_idx, main_sec in enumerate(self.syllabus.main_sections.all(), start=1):
            roman_index = int_to_roman(main_idx)
            sec_para = self.doc.add_paragraph()
            sec_para.paragraph_format.space_before = Pt(12)
            sec_para.paragraph_format.space_after = Pt(6)
            sec_run = sec_para.add_run(f"{roman_index}.\t{main_sec.name}")
            sec_run.font.name = 'Times New Roman'
            sec_run.font.bold = True
            sec_run.font.size = Pt(13)

            for sub_idx, sub_sec in enumerate(main_sec.sub_sections.all(), start=1):
                renderer = self.renderers.get(sub_sec.type)
                if renderer:
                    renderer.render(sub_sec, index=sub_idx)
        self._render_document_footer()
        buffer = io.BytesIO()
        self.doc.save(buffer)
        buffer.seek(0)
        return buffer

