import os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_sdt(parent, tag_name: str, placeholder: str = ""):
    sdt = OxmlElement('w:sdt')
    sdt_pr = OxmlElement('w:sdtPr')

    tag = OxmlElement('w:tag')
    tag.set(qn('w:val'), tag_name)
    sdt_pr.append(tag)

    alias = OxmlElement('w:alias')
    alias.set(qn('w:val'), tag_name)
    sdt_pr.append(alias)

    text_elem = OxmlElement('w:text')
    sdt_pr.append(text_elem)
    sdt.append(sdt_pr)

    sdt_content = OxmlElement('w:sdtContent')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = placeholder
    r.append(t)
    sdt_content.append(r)
    sdt.append(sdt_content)

    parent._p.append(sdt)
    return sdt


def wrap_table_in_sdt(table, tag_name: str):
    tbl_elem = table._tbl
    parent = tbl_elem.getparent()
    idx = parent.index(tbl_elem)

    sdt = OxmlElement('w:sdt')
    sdt_pr = OxmlElement('w:sdtPr')

    tag = OxmlElement('w:tag')
    tag.set(qn('w:val'), tag_name)
    sdt_pr.append(tag)

    alias = OxmlElement('w:alias')
    alias.set(qn('w:val'), tag_name)
    sdt_pr.append(alias)
    sdt.append(sdt_pr)

    sdt_content = OxmlElement('w:sdtContent')
    parent.remove(tbl_elem)
    sdt_content.append(tbl_elem)
    sdt.append(sdt_content)

    parent.insert(idx, sdt)


def build_template():
    output_dir = os.path.join("media", "templates", "docx")
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "template_master.docx")

    doc = Document()
    doc.add_heading("ĐỀ CƯƠNG CHI TIẾT HỌC PHẦN", level=1)

    p1 = doc.add_paragraph("Tên môn học: ")
    add_sdt(p1, "SUBJECT_NAME", "[Nhập tên môn học]")

    p2 = doc.add_paragraph("Mã môn học: ")
    add_sdt(p2, "SUBJECT_CODE", "[Nhập mã môn học]")

    p3 = doc.add_paragraph("Số tín chỉ: ")
    add_sdt(p3, "CREDIT", "[Nhập số tín chỉ]")

    doc.add_heading("Kế hoạch giảng dạy chi tiết", level=2)

    headers = [
        "Tuần",
        "Nội dung",
        "CĐR học phần",
        "Giờ tự học",
        "Giờ trực tiếp",
        "Giờ trực tuyến",
    ]
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"

    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    proto_data = [
        "[Tuần 1]",
        "[Nội dung bài học]",
        "[CLO 1, 2]",
        "[0]",
        "[0]",
        "[0]",
    ]
    for i, val in enumerate(proto_data):
        table.rows[1].cells[i].text = val

    wrap_table_in_sdt(table, "TEACHING_SCHEDULE")

    doc.save(file_path)
    print(f"File created successfully at: {os.path.abspath(file_path)}")


if __name__ == "__main__":
    build_template()